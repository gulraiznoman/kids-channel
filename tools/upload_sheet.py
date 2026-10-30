"""Write 'Episode N Upload Details.xlsx' (+ a .md copy) for one episode and its Shorts.

Used per episode, e.g. python tools/upload_sheet.py ep02
"""
import datetime as dt, pathlib, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parent.parent

EPISODES = {
    "ep02": dict(
        n=2, folder="02 Learn Colours", file="UPLOAD THIS - Learn Colours with Gul and Mishu v1.mp4",
        date=dt.date(2026, 10, 17), shorts_start=dt.date(2026, 10, 22),
        title="Learn Colours for Kids 🌈🎨 Colour Names for Toddlers | Gul & Mishu",
        intro="Let's learn colours with Gul & Mishu! 🌈🎨\n\nJoin Gul and Mishu in the rainbow garden as they discover 11 colours with "
              "big colourful balloons. Toddlers and preschoolers can see each colour, hear its name, and say it out loud. "
              "A fun first-words video for babies, toddlers and kids learning English.",
        list_title="Colours in this video:",
        closing="What colour are you wearing today? 👕\n\nGul & Mishu make fun learning videos for kids: fruits, colours, animals, numbers and more!",
        hashtags="#learncolors #learncolours #colorsforkids #kidslearning #toddlerlearning #preschool #firstwords #gulandmishu #educationalvideos",
        tags="learn colors, learn colours, colors for kids, colours for kids, color names, colors for toddlers, learn colors for toddlers, "
             "kids learning videos, toddler learning video, preschool learning, first words for babies, educational videos for kids, "
             "learn english for kids, gul and mishu, gul & mishu, kids cartoon",
        playlist_long="Learn with Gul & Mishu", playlist_short="Colours & Shapes",
        noun="Colour", topic="Learn Colours", full_name="Learn Colours for Kids",
        items=[("Red", "🔴", "Red, like an apple! Red is so bright!"),
               ("Orange", "🟠", "Orange, like an orange! That's funny!"),
               ("Yellow", "🟡", "Yellow, like a banana! Yellow like the sun!"),
               ("Green", "🟢", "Green, like a leaf! Green like the grass!"),
               ("Blue", "🔵", "Blue, like a little fish! Blue like the sky!"),
               ("Purple", "🟣", "Purple, like grapes! Purple is so pretty!"),
               ("Pink", "🩷", "Pink, like a flower! Pink like Mishu's dress!"),
               ("Brown", "🟤", "Brown, like a teddy bear! A big brown hug!"),
               ("Grey", "🩶", "Grey, like an elephant! A big grey elephant!"),
               ("Black", "⚫", "Black, like a kitty cat! Meow, meow!"),
               ("White", "⚪", "White, like a fluffy cloud! Soft and fluffy!")],
        tag_extra="color, colour, learn {low}, {low} color, {low} for kids, colors for toddlers, kids learning, preschool, gul and mishu",
    ),
}


def build(key):
    E = EPISODES[key]
    folder = ROOT / "VIDEOS" / E["folder"]
    desc = E["intro"] + "\n\n" + E["list_title"] + "\n" + "\n".join(f"{em} {nm}" for nm, em, _ in E["items"]) + \
        "\n\n" + E["closing"] + "\n\n" + E["hashtags"]
    rows = [[f"Episode {E['n']} (full video)", E["date"], E["file"], E["title"], desc, E["tags"],
             "THUMBNAIL option B.jpg", E["playlist_long"]]]
    for i, (nm, em, fact) in enumerate(E["items"], 1):
        low = nm.lower()
        rows.append([f"Short {i:02d}", E["shorts_start"] + dt.timedelta(days=i - 1), f"Shorts/Short {i:02d} - {nm}.mp4",
                     f"Guess the {E['noun']}! {em} {nm} | {E['topic']} for Toddlers #shorts",
                     f"Can you guess the {E['noun'].lower()}? {em} {fact}\nSay \"{low}\" out loud with Gul & Mishu!\n\n"
                     f"Watch the full video \"{E['full_name']}\" on the Gul & Mishu channel. 🌈\n\n"
                     f"#shorts #{low} #{E['topic'].lower().replace(' ', '')} #kidslearning #toddlers #preschool #gulandmishu",
                     E["tag_extra"].format(low=low), "No custom thumbnail: pick the frame with the big picture",
                     E["playlist_short"]])
    wb = Workbook()
    ws = wb.active
    ws.title = f"Episode {E['n']} uploads"
    hdr = ["Video", "Upload date", "File (in this folder)", "Title", "Description", "Tags", "Thumbnail", "Playlist",
           "Made for kids", "Category", "AI use", "Time (PKT)", "Status"]
    widths = [20, 12, 46, 52, 70, 50, 30, 20, 12, 12, 9, 10, 12]
    thin = Side(style="thin", color="C8CED9")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, (h, w) in enumerate(zip(hdr, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="2E5AAC")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = w
    dv = DataValidation(type="list", formula1='"Planned,Scheduled,Uploaded"', allow_blank=True)
    ws.add_data_validation(dv)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row + ["Yes", "Education", "No", "6:00 PM", "Planned"], 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name="Arial", size=10, bold=(r == 2 and c in (1, 4)))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = box
            if r == 2:
                cell.fill = PatternFill("solid", fgColor="FFF1C2")
        ws.cell(r, 2).number_format = "ddd dd mmm"
        ws.cell(r, 13).fill = PatternFill("solid", fgColor="FFF8DC")
        ws.cell(r, 13).font = Font(name="Arial", size=10, color="0000FF")
        dv.add(ws.cell(r, 13))
        ws.row_dimensions[r].height = 290 if r == 2 else 105
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    out = folder / f"Episode {E['n']} Upload Details.xlsx"
    wb.save(out)
    # markdown copy
    md = [f"# Episode {E['n']} upload details", "", "Same settings for every upload: Made for kids **Yes** · Category **Education** · AI use **No** · 6:00 PM PKT", ""]
    for row in rows:
        md += [f"## {row[0]} · {row[1].strftime('%a %d %b')}", f"File: `{row[2]}` · Playlist: {row[7]} · Thumbnail: {row[6]}", "",
               "**Title:**", "```", row[3], "```", "**Description:**", "```", row[4], "```", "**Tags:**", "```", row[5], "```", ""]
    (folder / f"Episode {E['n']} upload details.md").write_text("\n".join(md), encoding="utf-8")
    print("wrote", out.name, len(rows), "rows")


if __name__ == "__main__":
    build(sys.argv[1])
